#!/usr/bin/env python3
"""Смок: разделитель региона печатает фото раздела — ч/б и затемнённое (шаг 0151).

Решение владельца на тесте 24.09.2026: на разделителях регионов — одно фото,
ч/б и затемнённое, как фото раздела у эталона-ориентира
(`renderer/orion_golden_render/assets/divider-photo.webp`). Правила Р1–Р5 смок
проверяет на своём снимке — они о правиле, а не о конкретной картинке; Р6 — о
файле пакета: без него разделитель молча ушёл бы в полосы.

Проверяется:
  * Р1 — оба региона (Россия, ОАЭ) печатают панель фото, полос нет;
  * Р2 — фото ч/б и темнее исходного;
  * Р3 — панель в пределах полей и выше `INK_BOTTOM`: растровая проверка не
    терпит чернил за полями и ниже границы контента;
  * Р4 — заголовок держится левее панели в обоих вариантах разделителя;
  * Р5 — файла нет: полосы бренда, панели нет;
  * Р6 — фото пакета на месте и открывается, и это фото, а не миниатюра;
  * Р7 — у титула одна зелёная полоска, одна и та же в обоих вариантах, с фото
    и без: засечка заголовка и титул на её отступе. Замечание владельца с теста
    24.09.2026: у России (hero) были столб и серая засечка, у ОАЭ — одна полоска,
    «пусть как у ОАЭ».

Сеть и база не нужны; LibreOffice не нужен — проверяются фигуры PPTX.

Запуск: python3 renderer/smoke_divider_photo.py
"""

from __future__ import annotations

import io
import sys
import tempfile
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from PIL import Image, ImageStat  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

from smoke_counters import print_tap_counters  # noqa: E402
from deck_raster_layout import INK_BOTTOM  # noqa: E402
from orion_golden_render import slides as slides_module  # noqa: E402
from orion_golden_render.api import _draw_deck  # noqa: E402
from orion_golden_render.common import ACCENT, MARGIN_X, SLIDE_W  # noqa: E402

failures: list[str] = []
passed_checks = 0

#: Левее этого края начинается зона, где растровая проверка не терпит чернил.
SIDE_ZONE = int(MARGIN_X * 0.6)


def check(name: str, ok: bool, detail: str = "") -> None:
    global passed_checks
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if ok:
        passed_checks += 1
    else:
        failures.append(name)


def write_photo(path: Path) -> float:
    """Цветной снимок с градиентом; возвращает его среднюю яркость."""
    im = Image.new("RGB", (800, 600))
    px = im.load()
    for x in range(800):
        for y in range(600):
            px[x, y] = (60 + x * 180 // 800, 90 + y * 150 // 600, 200)
    im.save(path, "JPEG", quality=92)
    return ImageStat.Stat(im.convert("L")).mean[0]


def use_photo(path: Path) -> None:
    slides_module.DIVIDER_PHOTO_PATH = path
    slides_module._prepared_divider_photo.cache_clear()


def divider(section: str, title: str, variant: str | None = None) -> dict[str, Any]:
    slide: dict[str, Any] = {"template": "orion_golden_region_divider", "sectionKey": section, "title": title}
    if variant:
        slide["layoutVariant"] = variant
        slide["narrative"] = "Раздел отражает российское информационное поле, с которым столкнутся банки и партнёры."
    return slide


def draw(slides: list[dict[str, Any]]) -> Any:
    payload = {
        "deckManifest": {
            "finalSlides": [
                {**s, "slideKey": f"p{i:02d}", "pageNumber": i, "totalPageCount": len(slides)}
                for i, s in enumerate(slides, start=1)
            ]
        },
        "assets": [],
    }
    prs, *_rest = _draw_deck(payload, log_assets=False)
    return prs


def panel(prs: Any, page: int) -> Any:
    found = [sh for sh in prs.slides[page - 1].shapes if str(sh.name or "").startswith("orion_bg_divider")]
    return found[0] if found else None


def has_bands(prs: Any, page: int) -> bool:
    return any(str(sh.name or "").startswith("orion_decor_divider_band") for sh in prs.slides[page - 1].shapes)


def title_box(prs: Any, page: int) -> Any:
    boxes = [
        sh
        for sh in prs.slides[page - 1].shapes
        if getattr(sh, "has_text_frame", False) and "ЦИФРОВОЙ ПРОФИЛЬ" in sh.text_frame.text.upper()
    ]
    return boxes[0] if boxes else None


def stripes(prs: Any, page: int) -> list[Any]:
    """Вертикальные полоски листа: узкие автофигуры выше своей ширины.

    Полосы бренда и линия колонтитула — горизонтальные, панель фото — картинка:
    сюда они не попадают.
    """
    return [
        sh
        for sh in prs.slides[page - 1].shapes
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and int(sh.width) <= 200_000 and int(sh.height) > int(sh.width)
    ]


def stripe_signature(prs: Any, page: int) -> tuple[Any, ...] | str:
    """Полоска относительно титула: цвет, отступ, ширина, сдвиг от верха титула,
    высота на кегль титула и отступ самого титула. Кегли у регионов бывают разные
    (длинный титул уходит на ступень ниже), поэтому высота — в долях кегля."""
    found = stripes(prs, page)
    box = title_box(prs, page)
    if box is None:
        return "нет титула"
    if len(found) != 1:
        return f"полосок {len(found)}"
    bar = found[0]
    try:
        green = bar.fill.fore_color.rgb == ACCENT
    except Exception:  # noqa: BLE001
        green = False
    size = box.text_frame.paragraphs[0].runs[0].font.size
    return (
        "зелёная" if green else "не зелёная",
        int(bar.left) - MARGIN_X,
        int(bar.width),
        int(bar.top) - int(box.top),
        round(int(bar.height) / int(size), 2),
        int(box.left) - MARGIN_X,
    )


#: Файл фото пакета — до того, как Р1–Р5 подменят путь своим снимком.
PACKAGED_PHOTO = slides_module.DIVIDER_PHOTO_PATH


def main() -> int:
    try:
        with Image.open(PACKAGED_PHOTO) as packaged:
            size = packaged.size
        opened = True
    except Exception as exc:  # noqa: BLE001
        size, opened = (0, 0), False
        print(f"  фото пакета не открылось: {exc}")
    check(
        "Р6: фото пакета на месте, открывается и не миниатюра",
        opened and min(size) >= 400,
        f"{PACKAGED_PHOTO.name}: {size[0]}×{size[1]}" if opened else f"{PACKAGED_PHOTO}",
    )
    with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as tmp:
        photo = Path(tmp) / "divider-photo.jpg"
        source_mean = write_photo(photo)
        use_photo(photo)
        prs = draw(
            [
                divider("RU_PROFILE", "Россия: Цифровой профиль", "hero"),
                divider("UAE_PROFILE", "ОАЭ / международный: Цифровой профиль"),
            ]
        )
        ru, uae = panel(prs, 1), panel(prs, 2)
        check(
            "Р1: оба региона печатают панель фото, полос нет",
            ru is not None and uae is not None and not has_bands(prs, 1) and not has_bands(prs, 2),
            f"панель у России: {ru is not None}, у ОАЭ: {uae is not None}",
        )
        if ru is not None and ru.shape_type == MSO_SHAPE_TYPE.PICTURE:
            im = Image.open(io.BytesIO(ru.image.blob)).convert("RGB")
            r, g, b = (ImageStat.Stat(im).mean[i] for i in range(3))
            grey = max(abs(r - g), abs(g - b), abs(r - b)) < 2.0
            mean = ImageStat.Stat(im.convert("L")).mean[0]
            check(
                "Р2: фото ч/б и темнее исходного",
                grey and mean < source_mean * 0.8,
                f"каналы {r:.1f}/{g:.1f}/{b:.1f}, яркость {mean:.0f} при исходной {source_mean:.0f}",
            )
        else:
            check("Р2: фото ч/б и темнее исходного", False, "панели-картинки нет")
        if ru is not None:
            left, right = int(ru.left), int(ru.left) + int(ru.width)
            bottom = int(ru.top) + int(ru.height)
            check(
                "Р3: панель в пределах полей и выше INK_BOTTOM",
                left >= SIDE_ZONE and right <= SLIDE_W - SIDE_ZONE and bottom <= INK_BOTTOM and int(ru.top) >= 0,
                f"x {left}…{right} при зоне полей {SIDE_ZONE}…{SLIDE_W - SIDE_ZONE}, низ {bottom} при {INK_BOTTOM}",
            )
        else:
            check("Р3: панель в пределах полей и выше INK_BOTTOM", False, "панели нет")
        overlaps = []
        for page, pic in ((1, ru), (2, uae)):
            box = title_box(prs, page)
            if box is None or pic is None:
                overlaps.append(f"стр. {page}: нет заголовка или панели")
            elif int(box.left) + int(box.width) > int(pic.left):
                overlaps.append(f"стр. {page}: правый край заголовка {int(box.left) + int(box.width)} за левым краем панели {int(pic.left)}")
        check(
            "Р4: заголовок держится левее панели (hero и обычный разделитель)",
            not overlaps,
            "; ".join(overlaps) if overlaps else "обе страницы",
        )
        signatures = {"Россия (hero), с фото": stripe_signature(prs, 1), "ОАЭ, с фото": stripe_signature(prs, 2)}

        use_photo(Path(tmp) / "нет-такого-файла.jpg")
        prs = draw(
            [
                divider("RU_PROFILE", "Россия: Цифровой профиль", "hero"),
                divider("UAE_PROFILE", "ОАЭ / международный: Цифровой профиль"),
            ]
        )
        check(
            "Р5: файла нет — полосы бренда, панели нет (hero и обычный разделитель)",
            all(panel(prs, p) is None and has_bands(prs, p) for p in (1, 2)),
            "; ".join(f"стр. {p}: панель {panel(prs, p) is not None}, полосы {has_bands(prs, p)}" for p in (1, 2)),
        )
        signatures["Россия (hero), без фото"] = stripe_signature(prs, 1)
        signatures["ОАЭ, без фото"] = stripe_signature(prs, 2)
        reference = signatures["ОАЭ, с фото"]
        check(
            "Р7: у титула одна зелёная полоска, как у ОАЭ (hero и обычный, с фото и без)",
            isinstance(reference, tuple)
            and reference[0] == "зелёная"
            and all(sig == reference for sig in signatures.values()),
            "; ".join(f"{name}: {sig}" for name, sig in signatures.items()),
        )
    print(f"\n{'FAILED (' + str(len(failures)) + ')' if failures else 'PASSED (0 failures)'}")
    print_tap_counters(passed=passed_checks, failed=len(failures))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
