#!/usr/bin/env python3
"""Смок: строки блока доживают до листа и печатаются по ролям.

Приложение присылает блок списка строками (`\\n` — структурная строка: так режет
`clampClientText`, так раскладывает `reflowThemeBullet`, так велит писать модели
промпт стадии 2). Рендерер эти строки выбрасывал: `ctx.bullets` начинал с
`_safe(b)`, а тот схлопывает любые пробелы вместе с переводом строки. Дальше
структура угадывалась регулярками, знавшими июльскую форму текста
(`— источник домен`), и на сегодняшней (`— источник (адрес)`) не срабатывала —
клиент получал стену текста там, где приложение прислало шесть строк.

Проверяется оба конца:
  * сколько строк пришло, столько абзацев и нарисовано, и каждая оформлена по
    своей роли — заголовок, текст, цитата, мета, действие, адрес;
  * мера остаётся не оптимистичнее вывода, и судит её чужой прибор: страница
    конвертируется LibreOffice, и низ последней нарисованной строки сверяется с
    низом рамки, которую под список отвела мера.

Правило, которое здесь сторожится отдельно: **в цитате не выделяется ничего**.
Дословная цитата не редактируется даже весом шрифта.

Сеть и база не нужны. Нужны python-pptx, PyMuPDF и LibreOffice (для Т8).

Запуск: python3 renderer/smoke_line_typography.py
"""

from __future__ import annotations

import base64
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from smoke_counters import print_tap_counters  # noqa: E402
from orion_golden_render import common  # noqa: E402
from orion_golden_render.common import (  # noqa: E402
    BULLET_GLYPH,
    FS_BODY,
    FS_CAPTION,
    MUTED_COLOR,
    SLIDE_H,
    SLIDE_W,
    TONE_GOOD,
    _Ctx,
    _split_structured_bullet,
    get_bullet_measure,
    reset_bullet_measure,
    reset_layout_telemetry,
)

REPO_ROOT = Path(__file__).resolve().parent.parent

failures: list[str] = []
passed_checks = 0


def check(name: str, ok: bool, detail: str = "") -> None:
    global passed_checks
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if ok:
        passed_checks += 1
    else:
        failures.append(name)


#: Блок в том виде, в каком его сегодня отдаёт `toRendererPayload` золотого
#: кейса: строки уже есть, заголовок ещё приклеен к первой точкой (шаг 2 его
#: отделит), источник цитаты — полным адресом в скобках.
WIRE_BLOCK = "\n".join(
    [
        "Деловые связи, владение и контрагенты. Найдены конкретные материалы, в том числе "
        "«Профиль инвестора: семейные офисы Залива» (briefing-example.ae):",
        "«Глава фонда привлекает семейные офисы Залива» — источник "
        "(finans-example.se/glava-fonda-privlekaet-semejnye-ofisy-3)",
        "Деловые связи и круг связанных лиц могут расширять периметр проверки.",
    ]
)

#: Блок в том виде, в каком его отдаст построитель после шага 2: заголовок своей
#: строкой без точки, счётная фраза, источники, по строке на цитату, текст.
BUILDER_BLOCK = "\n".join(
    [
        "Назначения на руководящие должности",
        "По сюжету прочитано 8 публикаций, 1 из них нежелательная.",
        "Источники: site-a.example, site-b.example.",
        "«Планировалось, что коллегия рассмотрит 2 кандидатуры.» — источник "
        "(site-a.example/news/20195424.html).",
        "«На должность назначен указом на 6-летний срок» — источник (site-b.example/tag/page).",
        "Данные собраны 04.09.2026 по ТОП-20; ИНН 773800015809, решение 2019 года.",
        "Что делать: Проверить статусы 3 дел по картотекам судов.",
    ]
)

#: Лист «почему выделено» после шага 2: материал, основание, адрес.
WHY_BLOCK = "\n".join(
    [
        "«Фонд попал в список наблюдения» — kuriren-example.se",
        "Отнесено к теме «Сигналы списков наблюдения» по заголовку и описанию в выдаче; "
        "текст страницы в этом прогоне не проверялся.",
        "(kuriren-example.se/fond-popal-v-spisok-nablyudeniya-1).",
    ]
)

#: Плоский блок июльской формы — так написан замороженный текст эталона-72.
LEGACY_FLAT = (
    "«Деловой профиль» Найдены публикации по теме: «Биография основателя фонда» — источник "
    "highways-example.today «Предприниматель о трендах транспорта» — источник argumenti-example.ru. "
    "Где видно: highways-example.today, argumenti-example.ru."
)

#: Плоский блок сегодняшней формы: источник — адрес в скобках, домен кириллический.
TODAY_FLAT = (
    "«Судебные материалы» Найдены публикации по теме: «Суд назначил заседание по делу фонда» — "
    "источник (судьи-пример.рф/sudii/zasedanie-po-delu-fonda) «Фонд оспорил доначисление» — "
    "источник (kapital-example.se/fond-osporil-donachislenie-5). Всего по теме: 5 материалов, "
    "с негативным контекстом — 3."
)

#: Строка, склеенная на стороне приложения: цитата с адресом и мета в одной строке.
GLUED_BLOCK = "\n".join(
    [
        "«Офшорные структуры»",
        "Найдены публикации об офшорных структурах:",
        "«Фонд связан с холдингом на Мальте» — источник (watch-example.se/fond-svyazan-s-holdingom-2) "
        "Где видно: watch-example.se, reestr-example.ru.",
        "Всего по теме: 2 материала, с негативным контекстом — 2.",
    ]
)


def draw(items: list[str], *, page: int = 7) -> tuple[Any, Any]:
    """Нарисовать блоки на пустой странице; вернуть презентацию и рамку списка."""
    reset_layout_telemetry()
    reset_bullet_measure()
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    ctx = _Ctx(prs, page, 48, slide_key=f"p{page:02d}_typography")
    ctx.bullets(items, 1_230_000, max_items=9, max_chars=900)
    box = next(
        (
            sh
            for sh in prs.slides[0].shapes
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.startswith(BULLET_GLYPH)
        ),
        None,
    )
    return prs, box


def paragraphs(box: Any) -> list[Any]:
    return [p for p in box.text_frame.paragraphs if p.text.strip()] if box is not None else []


def text_runs(paragraph: Any) -> list[Any]:
    """Прогоны абзаца без маркера пункта."""
    return [r for r in paragraph.runs if r.text.strip() and not r.text.startswith(BULLET_GLYPH)]


def para_text(paragraph: Any) -> str:
    return "".join(r.text for r in text_runs(paragraph)).strip()


def bold_texts(paragraph: Any) -> list[str]:
    return [r.text.strip() for r in text_runs(paragraph) if r.font.bold]


def find(paras: list[Any], needle: str) -> Any | None:
    return next((p for p in paras if needle in para_text(p)), None)


def indent_of(paragraph: Any) -> tuple[int, int]:
    ppr = paragraph._p.pPr
    if ppr is None:
        return 0, 0
    return int(ppr.get("marL") or 0), int(ppr.get("indent") or 0)


def t1_lines_reach_the_page() -> None:
    _prs, box = draw([WIRE_BLOCK])
    paras = paragraphs(box)
    check(
        "Т1а: блок из трёх строк печатается тремя абзацами",
        len(paras) == 3,
        f"строк пришло 3, абзацев нарисовано {len(paras)}",
    )
    _prs, box = draw([BUILDER_BLOCK])
    paras = paragraphs(box)
    check(
        "Т1б: блок построителя из семи строк печатается семью абзацами",
        len(paras) == 7,
        f"абзацев нарисовано {len(paras)}",
    )
    drawn = " ".join(para_text(p) for p in paras)
    lost = [ln for ln in BUILDER_BLOCK.split("\n") if ln not in drawn]
    check("Т1в: ни одна строка блока не потеряна и не переписана", not lost, f"нет на листе: {lost[:2]}")

    # Темы резюме идут тем же путём списка, но текст к нему приходит через
    # дашборд — и там стоял свой `_safe`, склеивающий строки до `ctx.bullets`.
    from orion_golden_render.executive import _render_executive_dashboard

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    ctx = _Ctx(prs, 3, 48, slide_key="p03_executive")
    _render_executive_dashboard(
        ctx,
        {
            "narrative": "Итоговая оценка: высокий риск.",
            "metrics": [],
            # Блок построителя, а не сегодняшний: его семь строк угадыванием не
            # восстановить (заголовок без «ёлочек»), так что пройти проверку
            # можно только сохранив строки провода.
            "keyFindings": [{"detail": BUILDER_BLOCK, "tone": "warn"}],
            "actions": [],
        },
        "Резюме",
    )
    theme_box = next(
        (
            sh
            for sh in prs.slides[0].shapes
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.startswith(BULLET_GLYPH)
        ),
        None,
    )
    theme_paras = paragraphs(theme_box)
    check(
        "Т1г: тема резюме из семи строк печатается семью абзацами, заголовок — первым",
        len(theme_paras) == 7 and para_text(theme_paras[0]) == BUILDER_BLOCK.split("\n")[0],
        f"абзацев нарисовано {len(theme_paras)}",
    )


def t2_heading() -> None:
    _prs, box = draw([BUILDER_BLOCK])
    paras = paragraphs(box)
    head = paras[0] if paras else None
    check(
        "Т2а: заголовок блока — жирным целиком",
        head is not None and bool(text_runs(head)) and all(r.font.bold for r in text_runs(head)),
        f"жирные прогоны: {bold_texts(head) if head is not None else '—'}",
    )
    _prs, box = draw([WIRE_BLOCK])
    first = paragraphs(box)[0] if paragraphs(box) else None
    check(
        "Т2б: длинный ввод с двоеточием заголовком не становится",
        first is not None and not all(r.font.bold for r in text_runs(first)),
        f"жирные прогоны: {bold_texts(first) if first is not None else '—'}",
    )
    sentence_first = "Материалы по теме найдены в двух регионах.\nИсточники: site-a.example."
    _prs, box = draw([sentence_first])
    first = paragraphs(box)[0] if paragraphs(box) else None
    check(
        "Т2в: первая строка с конечной точкой — текст, а не заголовок",
        first is not None and not any(r.font.bold for r in text_runs(first)),
        f"жирные прогоны: {bold_texts(first) if first is not None else '—'}",
    )
    # Основание под заголовком обязано остаться обычным: на склеенном блоке старое
    # правило «строка с «ёлочки» — заголовок» красило жирным весь абзац целиком,
    # и проверка одной первой строки проходила по ложной причине.
    _prs, box = draw([WHY_BLOCK])
    why = paragraphs(box)
    first = why[0] if why else None
    basis = why[1] if len(why) > 1 else None
    check(
        "Т2г: «материал — источник» первой строкой листа «почему выделено» — заголовок, основание — нет",
        first is not None
        and basis is not None
        # Заголовок — строка провода дословно: старый разбор отрезал от неё
        # источник и уносил его в начало основания.
        and para_text(first) == WHY_BLOCK.split("\n")[0]
        and all(r.font.bold for r in text_runs(first))
        and not bold_texts(basis),
        f"абзацев {len(why)}; заголовок: {para_text(first)[:60] if first is not None else '—'}",
    )


def t3_quote() -> None:
    _prs, box = draw([BUILDER_BLOCK])
    quote = find(paragraphs(box), "Планировалось")
    runs = text_runs(quote) if quote is not None else []
    check(
        "Т3а: цитата печатается двумя прогонами — слова и атрибуция",
        len(runs) == 2 and runs[0].text.strip().startswith("«") and runs[1].text.strip().startswith("— источник"),
        f"прогоны: {[r.text[:28] for r in runs]}",
    )
    check(
        "Т3б: атрибуция цитаты — серым, слова цитаты — чернилами",
        len(runs) == 2
        and runs[1].font.color.rgb == MUTED_COLOR
        and runs[0].font.color.rgb != MUTED_COLOR,
        f"цвета: {[str(r.font.color.rgb) for r in runs]}",
    )
    check(
        "Т3в: внутри цитаты не выделено ничего, хотя в ней есть число",
        quote is not None and "2 кандидатуры" in para_text(quote) and not bold_texts(quote),
        f"жирные прогоны: {bold_texts(quote) if quote is not None else '—'}",
    )


def t4_numbers() -> None:
    _prs, box = draw([BUILDER_BLOCK])
    paras = paragraphs(box)
    stat = find(paras, "По сюжету прочитано")
    check(
        "Т4а: счётные числа строки выделены",
        stat is not None and bold_texts(stat) == ["8", "1"],
        f"жирные прогоны: {bold_texts(stat) if stat is not None else '—'}",
    )
    dates = find(paras, "Данные собраны")
    check(
        "Т4б: дата, «ТОП-20», ИНН и год числами не считаются",
        dates is not None and not bold_texts(dates),
        f"жирные прогоны: {bold_texts(dates) if dates is not None else '—'}",
    )
    steps = "Следующие проверки\n1) Сверить первоисточники по 5 темам.\n2) Сверить записи (1) базы."
    _prs, box = draw([steps])
    first_step = find(paragraphs(box), "Сверить первоисточники")
    check(
        "Т4в: номер пункта выделен вместе со скобкой и только в начале строки",
        first_step is not None and bold_texts(first_step) == ["1)", "5"],
        f"жирные прогоны: {bold_texts(first_step) if first_step is not None else '—'}",
    )


def t4_foreign_text_is_not_emphasized() -> None:
    """Выделяются наши законченные предложения; чужой текст не трогается.

    Чужие слова приходят в блок четырьмя путями: в «ёлочках», после ярлыка
    «(дословно):», строкой без конечного знака (подсказка, связанный запрос,
    заголовок) и ответом поискового ИИ. Ни в одном из них рендерер не вправе
    расставлять акценты.
    """
    block = "\n".join(
        [
            "Сводка по теме",
            "Примеры: Фонд выделил 16 миллионов евро на сделку · Суд рассмотрит 3 иска.",
            "Всего по теме: 5 материалов, с негативным контекстом — 3.",
        ]
    )
    _prs, box = draw([block])
    paras = paragraphs(box)
    examples = find(paras, "Примеры:")
    totals = find(paras, "Всего по теме:")
    check(
        "Т4г: в строке «Примеры:» — чужие заголовки, числа в них не выделены",
        examples is not None and bold_texts(examples) == ["Примеры:"],
        f"жирные прогоны: {bold_texts(examples) if examples is not None else '—'}",
    )
    check(
        "Т4д: в счётной строке «Всего по теме:» числа выделены",
        totals is not None and bold_texts(totals) == ["Всего по теме:", "5", "3"],
        f"жирные прогоны: {bold_texts(totals) if totals is not None else '—'}",
    )
    _prs, box = draw(["Фонд 5 причин банкротства", "Фонд: биография основателя"])
    paras = paragraphs(box)
    check(
        "Т4е: строка без конечного знака — чужая (подсказка, запрос): ни чисел, ни ярлыка",
        len(paras) == 2 and not bold_texts(paras[0]) and not bold_texts(paras[1]),
        f"жирные прогоны: {[bold_texts(p) for p in paras]}",
    )
    _prs, box = draw(["Начало статьи (дословно): Основал 3 компании в 2 странах."])
    verbatim = paragraphs(box)[0] if paragraphs(box) else None
    check(
        "Т4ж: после ярлыка «(дословно):» числа не выделяются — ярлык наш, слова чужие",
        verbatim is not None and bold_texts(verbatim) == ["Начало статьи (дословно):"],
        f"жирные прогоны: {bold_texts(verbatim) if verbatim is not None else '—'}",
    )

    # Ответ поискового ИИ — чужой текст без кавычек; страницу узнаём по шаблону
    # деки, в том числе когда картинки нет и она идёт прозаическим макетом.
    from orion_golden_render.slides import _render_slide

    answer = "Ответ поискового ИИ Яндекса. Запрос: «Фонд». Фонд управляет 12 компаниями в 3 странах."
    for template in ("orion_golden_prose", "orion_golden_surface_panel"):
        prs = Presentation()
        prs.slide_width = Emu(SLIDE_W)
        prs.slide_height = Emu(SLIDE_H)
        ctx = _Ctx(prs, 50, 80, slide_key="p50_ai")
        _render_slide(
            ctx,
            {
                "template": template,
                "templateId": "ai-overview",
                "isContinuation": True,
                "title": "Россия — AI-ответы поисковых систем",
                "bullets": [answer],
            },
            {},
        )
        ai_box = next(
            (
                sh
                for sh in prs.slides[0].shapes
                if getattr(sh, "has_text_frame", False) and sh.text_frame.text.startswith(BULLET_GLYPH)
            ),
            None,
        )
        ai_paras = paragraphs(ai_box)
        check(
            f"Т4з: на странице AI-ответов ({template.replace('orion_golden_', '')}) в ответе ничего не выделено",
            bool(ai_paras) and not any(bold_texts(p) for p in ai_paras),
            f"жирные прогоны: {[bold_texts(p) for p in ai_paras]}",
        )


def t5_meta_and_action() -> None:
    _prs, box = draw([BUILDER_BLOCK])
    paras = paragraphs(box)
    meta = find(paras, "Источники:")
    meta_runs = text_runs(meta) if meta is not None else []
    check(
        "Т5а: мета-строка — кеглем подписи, серым, ярлык жирным",
        bool(meta_runs)
        and all(r.font.size.pt == FS_CAPTION and r.font.color.rgb == MUTED_COLOR for r in meta_runs)
        and bold_texts(meta) == ["Источники:"],
        f"кегли {[r.font.size.pt for r in meta_runs]}, жирные {bold_texts(meta) if meta is not None else '—'}",
    )
    action = find(paras, "Что делать:")
    action_runs = text_runs(action) if action is not None else []
    check(
        "Т5б: ярлык действия — жирным тоном действия, текст — основным кеглем",
        bool(action_runs)
        and action_runs[0].text.strip() == "Что делать:"
        and action_runs[0].font.bold
        and action_runs[0].font.color.rgb == TONE_GOOD
        and all(r.font.size.pt == FS_BODY for r in action_runs),
        f"первый прогон: {action_runs[0].text if action_runs else '—'}",
    )
    _prs, box = draw([WHY_BLOCK])
    address = find(paragraphs(box), "(kuriren-example.se/")
    address_runs = text_runs(address) if address is not None else []
    check(
        "Т5в: строка-адрес — подписью серым, без выделений",
        bool(address_runs)
        and all(r.font.size.pt == FS_CAPTION and r.font.color.rgb == MUTED_COLOR for r in address_runs)
        and not bold_texts(address),
        f"кегли {[r.font.size.pt for r in address_runs]}",
    )


def t6_hanging_indent() -> None:
    _prs, box = draw([BUILDER_BLOCK])
    paras = paragraphs(box)
    head_mar, head_indent = indent_of(paras[0]) if paras else (0, 0)
    check(
        "Т6а: у первой строки блока висячий отступ — перенос встаёт под текст",
        head_mar > 0 and head_indent == -head_mar,
        f"marL={head_mar}, indent={head_indent}",
    )
    text_para = find(paras, "По сюжету прочитано")
    quote_para = find(paras, "Планировалось")
    text_mar = indent_of(text_para)[0] if text_para is not None else 0
    quote_mar = indent_of(quote_para)[0] if quote_para is not None else 0
    check(
        "Т6б: строки блока стоят под текстом заголовка, цитата — глубже",
        text_mar == head_mar and quote_mar > text_mar > 0,
        f"заголовок {head_mar}, текст {text_mar}, цитата {quote_mar}",
    )
    check(
        "Т6в: отступ пробелами больше не печатается",
        all(not r.text.startswith("   ") for p in paras for r in p.runs),
        "",
    )


def t7_flat_input_still_splits() -> None:
    legacy = _split_structured_bullet(common._safe_preserve_breaks(LEGACY_FLAT))
    check(
        "Т7а: плоский блок июльской формы раскладывается как раньше",
        len(legacy) == 5 and legacy[0] == "«Деловой профиль»" and legacy[-1].startswith("Где видно:"),
        f"строк {len(legacy)}: {[ln[:24] for ln in legacy]}",
    )
    today = _split_structured_bullet(common._safe_preserve_breaks(TODAY_FLAT))
    quotes = [ln for ln in today if ln.startswith("«") and "— источник (" in ln]
    check(
        "Т7б: плоский блок с источником-адресом раскладывается по цитатам",
        len(quotes) == 2 and any("судьи-пример.рф" in q for q in quotes),
        f"строк {len(today)}, цитат своей строкой {len(quotes)}",
    )
    glued = _split_structured_bullet(common._safe_preserve_breaks(GLUED_BLOCK))
    check(
        "Т7в: мета, приклеенная к цитате с адресом, уходит на свою строку",
        any(ln.startswith("Где видно:") for ln in glued)
        and not any("Где видно:" in ln and ln.startswith("«") for ln in glued),
        f"строки: {[ln[:32] for ln in glued]}",
    )
    between = (
        "«Тема с пояснением» Найдены публикации по теме: «Первая цитата о фонде и его сделке» — "
        "источник a-example.ru Пояснение между цитатами называет сумму сделки. «Вторая цитата о "
        "фонде и его сделке» — источник (b-example.ru/sdelka-fonda-2). Всего по теме: 2 материала."
    )
    kept = _split_structured_bullet(common._safe_preserve_breaks(between))
    check(
        "Т7д: слова между цитатами при пересборке не пропадают",
        any("Пояснение между цитатами называет сумму сделки" in ln for ln in kept),
        f"строки: {[ln[:36] for ln in kept]}",
    )
    _prs, box = draw([LEGACY_FLAT])
    head = paragraphs(box)[0] if paragraphs(box) else None
    check(
        "Т7г: тема в «ёлочках» по-прежнему печатается заголовком",
        head is not None and para_text(head) == "«Деловой профиль»" and all(r.font.bold for r in text_runs(head)),
        f"первая строка: {para_text(head) if head is not None else '—'}",
    )


def _heavy_pages() -> list[list[str]]:
    long_quote = (
        "«Коллегия рассмотрит кандидатуры на должность председателя суда и заместителей, а также "
        "вопрос о продлении полномочий действующего состава до конца следующего года» — источник "
        "(site-a.example/news/kollegiya-rassmotrit-kandidatury-na-dolzhnost-predsedatelya-20195424.html)."
    )
    dense = "\n".join(
        ["Плотная тема", "По сюжету прочитано 12 публикаций, 4 из них нежелательные."]
        + [f"Источники: site-{i}.example, portal-{i}.example, agency-{i}.example." for i in range(3)]
        + ["Что делать: Сверить статусы дел и первоисточники до принятия решений."]
    )
    quoted = "\n".join(
        ["Тема с длинными цитатами", "По сюжету прочитано 9 публикаций, 9 из них нежелательные."]
        + [long_quote] * 3
        + ["Сведения требуют проверки по первичным документам."]
    )
    wrapped_heading = "\n".join(
        [
            "Очень длинный заголовок блока о назначениях, выдвижениях и переназначениях на руководящие "
            "должности в судах региона и смежных ведомствах",
            "По сюжету прочитано 3 публикации.",
        ]
    )
    return [
        [dense] * 5,
        [quoted] * 3,
        [wrapped_heading] * 8,
        [BUILDER_BLOCK, WHY_BLOCK, dense, quoted],
    ]


def t8_measure_is_not_optimistic() -> None:
    """Мера против чужого прибора: вёрстка LibreOffice, строки из PDF."""
    try:
        import fitz  # noqa: F401
        from orion_golden_render.api import measure_orion_golden, render_orion_golden
    except Exception as exc:  # noqa: BLE001
        print(f"# SKIP Т8 мера против отрисовки — нет зависимостей рендера: {exc}")
        return
    import fitz

    pages = _heavy_pages()
    payload: dict[str, Any] = {
        "reportSpec": {"subject": {"displayName": "Субъект Проверки"}},
        "deckManifest": {
            "finalSlides": [
                {
                    "slideKey": f"p{i:02d}_prose",
                    "template": "orion_golden_prose",
                    "title": f"Материалы повышенного внимания — лист {i}",
                    "pageNumber": i,
                    "totalPageCount": len(pages),
                    "bullets": list(bullets),
                }
                for i, bullets in enumerate(pages, start=1)
            ]
        },
        "assets": [],
    }
    verdict = {p["slideKey"]: p for p in measure_orion_golden(payload)["pages"]}
    # Лист набирается по мере — тем же правилом, что `planBulletRecut`.
    for slide in payload["deckManifest"]["finalSlides"]:
        page = verdict[slide["slideKey"]]
        used = fit = 0
        for h in page["itemHeights"][: page["maxItems"]]:
            if fit > 0 and used + h > page["availableHeight"]:
                break
            used += h
            fit += 1
        slide["bullets"] = slide["bullets"][:fit]
    out = render_orion_golden(payload)
    if out.get("pdfExportMode") != "libreoffice":
        print("# SKIP Т8 мера против отрисовки — LibreOffice недоступен, вёрстку судить нечем")
        return
    lost = [
        e
        for e in (out.get("layoutTelemetry") or {}).get("entries") or []
        if e.get("droppedBullets") or e.get("droppedLines")
    ]
    check("Т8а: набранные по мере листы рисуются без потерь", not lost, f"потери: {[e.get('name') for e in lost]}")

    prs = Presentation(__import__("io").BytesIO(base64.b64decode(out["pptxBase64"])))
    doc = fitz.open(stream=base64.b64decode(out["pdfBase64"]), filetype="pdf")
    worst: tuple[int, float] | None = None
    checked = 0
    for index, slide in enumerate(prs.slides):
        box = next(
            (
                sh
                for sh in slide.shapes
                if getattr(sh, "has_text_frame", False) and sh.text_frame.text.startswith(BULLET_GLYPH)
            ),
            None,
        )
        if box is None:
            continue
        box_top_pt = box.top / 12_700
        box_bottom_pt = (box.top + box.height) / 12_700
        footer_top_pt = (SLIDE_H - 480_000) / 12_700
        bottoms = [
            span["bbox"][3]
            for block in doc[index].get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line.get("spans", [])
            if span["text"].strip() and box_top_pt - 2 <= span["bbox"][1] < footer_top_pt
        ]
        if not bottoms:
            continue
        checked += 1
        spill = max(bottoms) - box_bottom_pt
        if worst is None or spill > worst[1]:
            worst = (index + 1, spill)
    check("Т8б: проверка не вакуумна — листы со списком найдены в PDF", checked == len(pages), f"листов {checked}")
    check(
        "Т8в: последняя нарисованная строка не ниже рамки, отведённой мерой",
        worst is not None and worst[1] <= 0.5,
        f"худший лист {worst[0]}: строка ниже рамки на {worst[1]:.1f} pt" if worst else "нечего мерить",
    )


def t9_one_vocabulary() -> None:
    src_json = REPO_ROOT / "src/modules/digital-profile/orion-golden/client/client-text-contract.json"
    renderer_json = REPO_ROOT / "renderer/client_text_contract.json"
    same = (
        hashlib.sha256(src_json.read_bytes()).hexdigest()
        == hashlib.sha256(renderer_json.read_bytes()).hexdigest()
    )
    check("Т9а: копии контракта побайтно равны", same, "")
    contract = json.loads(renderer_json.read_text(encoding="utf-8"))
    section = contract.get("typography") or {}
    check(
        "Т9б: словарь форм объявлен в контракте",
        bool(section.get("metaLabels")) and bool(section.get("actionLabels")),
        f"ключи раздела: {sorted(section)}",
    )
    labels = list(section.get("metaLabels") or [])
    missed = [label for label in labels if not common._META_LINE_RE.match(f"{label}: пример")]
    check(
        "Т9в: `_META_LINE_RE` собран из того же списка, что и роли",
        bool(labels) and not missed,
        f"не узнаны: {missed}" if labels else "списка нет",
    )
    added = "Сводка по региону"
    try:
        from orion_golden_render import typography
    except ImportError:
        check("Т9г: новый ярлык контракта становится мета-строкой без правки кода", False, "модуля типографики нет")
        return
    patched = dict(contract)
    patched["typography"] = {**section, "metaLabels": [*labels, added]}
    layout = typography.line_layout(f"{added}: 3 темы.", index=1, total=3, contract=patched)
    check(
        "Т9г: новый ярлык контракта становится мета-строкой без правки кода",
        layout.role == typography.ROLE_META,
        f"роль: {layout.role}",
    )


# --------------------------------------------------------------------------
# Абзац страницы (шаг 0098)
# --------------------------------------------------------------------------

#: Абзац резюме в том виде, в каком его отдаст композитор: оценка, перечень,
#: подзаголовки своими строками, основания по строке.
NARRATIVE = "\n".join(
    [
        "Итоговая оценка: высокий риск.",
        "Основные основания: Деловые связи и контрагенты; Судебные материалы; Офшоры.",
        "Главные основания",
        "Деловые связи и контрагенты: фонд привлекает семейные офисы Залива.",
        "Судебные материалы: основателю фонда предъявлено налоговое обвинение.",
        "Ограничения",
        "Вывод основан на открытых источниках; первичные документы могут изменить оценку.",
        "Исследованы результаты поиска (ТОП-20) по 2 регионам. Данные собраны 04.09.2026.",
    ]
)


def draw_body(text: str, **kwargs: Any) -> tuple[Any, Any, list[dict[str, Any]]]:
    """Нарисовать абзац на пустой странице; вернуть презентацию, рамку и телеметрию."""
    from orion_golden_render.common import get_layout_telemetry

    reset_layout_telemetry()
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    ctx = _Ctx(prs, 4, 48, slide_key="p04_body")
    ctx.body(text, 1_230_000, **kwargs)
    box = next(
        (sh for sh in prs.slides[0].shapes if getattr(sh, "name", "").startswith("orion_text_body")),
        None,
    )
    return prs, box, get_layout_telemetry()


def body_runs(paragraph: Any) -> list[Any]:
    return [r for r in paragraph.runs if r.text.strip()]


def body_bold(paragraph: Any) -> list[str]:
    return [r.text.strip() for r in body_runs(paragraph) if r.font.bold]


def a1_a4_paragraphs_and_lead() -> None:
    lines = NARRATIVE.split("\n")
    _prs, box, _tel = draw_body(NARRATIVE, max_h=4_000_000, bold=True)
    paras = paragraphs(box)
    check(
        "А1: абзацы провода рисуются абзацами",
        len(paras) == len(lines),
        f"пришло {len(lines)}, нарисовано {len(paras)}",
    )
    lead = paras[0] if paras else None
    rest_bold = [body_bold(p) for p in paras[1:]]
    check(
        "А2: при bold=True жирный целиком только первый абзац — вывод",
        lead is not None
        and all(r.font.bold for r in body_runs(lead))
        and not any(p.text.strip() and all(r.font.bold for r in body_runs(p)) and p.text.strip()[-1:] in ".!?…" for p in paras[1:]),
        f"жирное дальше первого абзаца: {rest_bold[:4]}",
    )
    heads = [p for p in paras if p.text.strip() in ("Главные основания", "Ограничения")]
    check(
        "А3а: абзац без конечного знака — подзаголовок, жирным целиком",
        len(heads) == 2 and all(all(r.font.bold for r in body_runs(p)) for p in heads),
        f"найдено подзаголовков {len(heads)}",
    )
    _prs, box, _tel = draw_body("Короткая строка без точки", max_h=900_000)
    single = paragraphs(box)
    check(
        "А3б: единственный абзац подзаголовком не становится",
        len(single) == 1 and not body_bold(single[0]),
        f"жирное: {body_bold(single[0]) if single else '—'}",
    )
    _prs, box, _tel = draw_body(NARRATIVE, max_h=4_000_000, bold=True)
    grounds = next((p for p in paragraphs(box) if p.text.strip().startswith("Основные основания:")), None)
    scope = next((p for p in paragraphs(box) if p.text.strip().startswith("Исследованы")), None)
    check(
        "А4а: ярлык абзаца и числа выделены; дата и «ТОП-20» — нет",
        grounds is not None
        and body_bold(grounds) == ["Основные основания:"]
        and scope is not None
        and body_bold(scope) == ["2"],
        f"ярлык: {body_bold(grounds) if grounds is not None else '—'}; числа: {body_bold(scope) if scope is not None else '—'}",
    )
    caption = "Доля негатива: 17 % (10 из 58).\nСтраницы о других людях (1) в долю не входят."
    _prs, box, _tel = draw_body(caption, max_h=900_000, color=MUTED_COLOR, font_size=FS_CAPTION)
    check(
        "А4б: в подписи серым ничего не выделяется",
        bool(paragraphs(box)) and not any(body_bold(p) for p in paragraphs(box)),
        f"жирное: {[body_bold(p) for p in paragraphs(box)]}",
    )


def a5_single_paragraph_measures_as_before() -> None:
    """Один абзац — та же формула: на ней держится ёмкость страниц выдачи."""
    from orion_golden_render.common import CONTENT_W, measure_text_height

    intro = (
        "Показана выдача Яндекса по запросу «Фонд Северный капитал»: 20 позиций, из них 3 отмечены "
        "как нежелательные. Страница формирует первое впечатление о субъекте у банка и партнёра. "
        "Проверить первоисточники выделенных результатов и сверить статусы дел."
    )
    for label, kwargs in (
        ("обычным", {}),
        ("жирным", {"bold": True}),
        ("подписью серым", {"color": MUTED_COLOR}),
    ):
        _prs, _box, tel = draw_body(intro, max_h=1_000_000, **kwargs)
        entry = next((e for e in tel if str(e.get("name", "")).startswith("orion_text_body")), None)
        expected = measure_text_height(
            intro, CONTENT_W, FS_BODY, line_spacing=1.2, paragraph_spacing_pt=8, bold=bool(kwargs.get("bold"))
        )
        # Жирный лид и подпись серым меряются ровно как прежде — на подписи серым
        # держится ёмкость страниц выдачи. Основной текст вправе стать
        # консервативнее: строка с выделенным числом меряется жирной целиком.
        exact = label != "обычным"
        check(
            f"А5: один абзац {label} меряется {'той же формулой, что до шага' if exact else 'не оптимистичнее, чем до шага'}",
            entry is not None
            and (entry["requiredHeight"] == expected if exact else entry["requiredHeight"] >= expected),
            f"мера {entry['requiredHeight'] if entry else '—'}, прежняя формула {expected}",
        )


def a5b_bold_lead_is_measured_bold() -> None:
    """Жирный лид меряется жирным — на ширине, где это видно.

    На ширине колонки жирный и обычный текст переносятся одинаково, и сверка
    высот их не различает: мутация «мерить абзац обычным начертанием» оставалась
    зелёной. Ширина ищется та, на которой лишние проценты жирного дают лишний
    перенос, — тот же приём, что в `smoke_text_measurement.py`.
    """
    from orion_golden_render.common import measure_text_height

    lead = "Итоговая оценка по открытым источникам: высокий риск для деловой репутации проверяемого лица."
    found = next(
        (
            w
            for w in range(1_500_000, 9_000_001, 50_000)
            if measure_text_height(lead, w, FS_BODY, line_spacing=1.2, bold=True)
            > measure_text_height(lead, w, FS_BODY, line_spacing=1.2, bold=False)
        ),
        None,
    )
    if found is None:
        check("А5б: найдена ширина, на которой жирный лид выше обычного", False, "не найдена")
        return
    _prs, _box, tel = draw_body(lead + "\nВторой абзац обычным начертанием.", max_h=3_000_000, bold=True, w=found)
    entry = next((e for e in tel if str(e.get("name", "")).startswith("orion_text_body")), None)
    second = measure_text_height("Второй абзац обычным начертанием.", found, FS_BODY, line_spacing=1.2, bold=False)
    bold_h = measure_text_height(lead, found, FS_BODY, line_spacing=1.2, bold=True)
    regular_h = measure_text_height(lead, found, FS_BODY, line_spacing=1.2, bold=False)
    gap = int(8 * 12_700 * 1.18)
    check(
        "А5б: жирный лид меряется жирным, абзац под ним — обычным",
        entry is not None and entry["requiredHeight"] == bold_h + second + gap and bold_h > regular_h,
        f"мера {entry['requiredHeight'] if entry else '—'}; жирным {bold_h + second + gap}, обычным {regular_h + second + gap} (ширина {found})",
    )


def a6_a7_nothing_is_dropped_silently() -> None:
    nine = "\n".join(f"Абзац номер {i}: проверка того, что срез по числу абзацев снят." for i in range(1, 10))
    _prs, box, _tel = draw_body(nine, max_h=4_500_000)
    check(
        "А6: девять абзацев рисуются все девять — молчаливого среза по числу нет",
        len(paragraphs(box)) == 9,
        f"нарисовано {len(paragraphs(box))}",
    )
    many = "\n".join(
        f"Абзац {i}: " + "длинное предложение о проверке первоисточников и статусов дел. " * 4
        for i in range(1, 9)
    )
    _prs, box, tel = draw_body(many, max_h=1_200_000)
    drawn = paragraphs(box)
    entry = next((e for e in tel if str(e.get("name", "")).startswith("orion_text_body")), None)
    check(
        "А7а: не влезающий текст — абзацы сняты с конца, первый на месте",
        0 < len(drawn) < 8 and drawn[0].text.strip().startswith("Абзац 1:"),
        f"нарисовано {len(drawn)} из 8",
    )
    check(
        "А7б: срез слышен — запись телеметрии помечена клипом",
        entry is not None and entry.get("clipped") is True,
        f"clipped={entry.get('clipped') if entry else '—'}",
    )
    whole = "Короткий абзац целиком.\nВторой короткий абзац."
    _prs, _box, tel = draw_body(whole, max_h=1_200_000)
    entry = next((e for e in tel if str(e.get("name", "")).startswith("orion_text_body")), None)
    check(
        "А7в: влезший целиком текст клипом не помечен",
        entry is not None and entry.get("clipped") is False,
        f"clipped={entry.get('clipped') if entry else '—'}",
    )


def a8_search_table_intro_is_one_paragraph() -> None:
    from orion_golden_render.slides import _render_slide

    intro = "Показана выдача Яндекса: 20 позиций.\nСтраница формирует первое впечатление.\nПроверить первоисточники."
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    ctx = _Ctx(prs, 16, 80, slide_key="p09_ru_serp_table")
    _render_slide(
        ctx,
        {
            "template": "orion_golden_search_table",
            "title": "Россия — Яндекс: собранная выдача",
            "narrative": intro,
            "table": {"headers": ["№", "Заголовок", "Оценка"], "rows": [["1", "Материал", "Нейтральный"]]},
        },
        {},
    )
    box = next(
        (sh for sh in prs.slides[0].shapes if getattr(sh, "name", "").startswith("orion_text_body")),
        None,
    )
    check(
        "А8: вводный абзац страницы выдачи остаётся одним абзацем — его ёмкость откалибрована так",
        len(paragraphs(box)) == 1,
        f"абзацев {len(paragraphs(box))}",
    )


def a9_body_measure_is_not_optimistic() -> None:
    """Мера многоабзацного текста против вёрстки LibreOffice."""
    try:
        import fitz  # noqa: F401
        from orion_golden_render.api import render_orion_golden
    except Exception as exc:  # noqa: BLE001
        print(f"# SKIP А9 мера абзаца против отрисовки — нет зависимостей рендера: {exc}")
        return
    import fitz

    long_lines = [
        "Итоговая оценка: высокий риск.",
        "Основные основания: " + "; ".join(f"тема риска номер {i} с длинным названием" for i in range(1, 7)) + ".",
        "Главные основания",
    ] + [
        f"Тема риска номер {i}: публикация о споре вокруг актива фонда и его контрагентов в двух юрисдикциях, "
        f"повторённая {i + 2} изданиями в течение квартала."
        for i in range(1, 6)
    ] + ["Ограничения", "Вывод основан на открытых источниках; первичные документы могут изменить оценку."]
    payload = {
        "reportSpec": {"subject": {"displayName": "Субъект Проверки"}},
        "deckManifest": {
            "finalSlides": [
                {
                    "slideKey": "p03_executive",
                    "template": "orion_golden_executive_dashboard",
                    "title": "Резюме",
                    "pageNumber": 1,
                    "totalPageCount": 1,
                    "narrative": "\n".join(long_lines),
                    "metrics": [{"label": "Материалов собрано", "value": "341", "tone": "neutral"}],
                    "keyFindings": [],
                }
            ]
        },
        "assets": [],
    }
    out = render_orion_golden(payload)
    if out.get("pdfExportMode") != "libreoffice":
        print("# SKIP А9 мера абзаца против отрисовки — LibreOffice недоступен")
        return
    prs = Presentation(__import__("io").BytesIO(base64.b64decode(out["pptxBase64"])))
    doc = fitz.open(stream=base64.b64decode(out["pdfBase64"]), filetype="pdf")
    box = next(
        (sh for sh in prs.slides[0].shapes if getattr(sh, "name", "").startswith("orion_text_body")),
        None,
    )
    if box is None:
        check("А9: абзац резюме нарисован", False, "рамки абзаца нет")
        return
    top_pt, bottom_pt = box.top / 12_700, (box.top + box.height) / 12_700
    footer_top_pt = (SLIDE_H - 480_000) / 12_700
    bottoms = [
        span["bbox"][3]
        for block in doc[0].get_text("dict")["blocks"]
        for line in block.get("lines", [])
        for span in line.get("spans", [])
        if span["text"].strip() and top_pt - 2 <= span["bbox"][1] < footer_top_pt
    ]
    drawn = len(paragraphs(box))
    check("А9а: абзацы резюме нарисованы одним блоком", drawn == len(long_lines), f"абзацев {drawn} из {len(long_lines)}")
    check(
        "А9б: последняя строка абзаца не ниже рамки, отведённой мерой",
        bool(bottoms) and max(bottoms) - bottom_pt <= 0.5,
        f"строка ниже рамки на {max(bottoms) - bottom_pt:.1f} pt" if bottoms else "строк не найдено",
    )


def main() -> int:
    t1_lines_reach_the_page()
    t2_heading()
    t3_quote()
    t4_numbers()
    t4_foreign_text_is_not_emphasized()
    t5_meta_and_action()
    t6_hanging_indent()
    t7_flat_input_still_splits()
    t8_measure_is_not_optimistic()
    t9_one_vocabulary()
    a1_a4_paragraphs_and_lead()
    a5_single_paragraph_measures_as_before()
    a5b_bold_lead_is_measured_bold()
    a6_a7_nothing_is_dropped_silently()
    a8_search_table_intro_is_one_paragraph()
    a9_body_measure_is_not_optimistic()

    print(f"\n{'FAILED (' + str(len(failures)) + ')' if failures else 'PASSED (0 failures)'}")
    print_tap_counters(passed=passed_checks, failed=len(failures))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
