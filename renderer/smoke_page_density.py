#!/usr/bin/env python3
"""Смок: плотность листа — полоса заголовка по факту и нижний запас (шаг 0114).

Рендерер клал содержимое от `y + 950 000` после заголовка независимо от числа
его строк и держал нижнюю границу на `SLIDE_H − 1 100 000` при подвале на
`SLIDE_H − 440 000`: у большинства страниц заголовок однострочный, и над
подвалом пустовало 660 000 EMU. Живые отчёты 18.09.2026 — страницы с одним-двумя
блоками и пустой нижней третью.

Проверяется:
  * П1 — однострочный заголовок отдаёт `y + 700 000`, двухстрочный — `y + 950 000`;
  * П2 — `CONTENT_BOTTOM == SLIDE_H − 800 000`;
  * П3 — список под однострочным заголовком получает в мере на 550 000 больше
    прежнего листа (4 865 200 → 5 415 200);
  * П4 — подвал не задет: низ рамки списка не ниже `CONTENT_BOTTOM`, подвал ниже.

Сеть и база не нужны. Нужен python-pptx.

Запуск: python3 renderer/smoke_page_density.py
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from smoke_counters import print_tap_counters  # noqa: E402
from orion_golden_render.common import (  # noqa: E402
    BULLET_GLYPH,
    CONTENT_BOTTOM,
    FOOTER_Y,
    SLIDE_H,
    SLIDE_W,
    _Ctx,
    get_bullet_measure,
    reset_bullet_measure,
    reset_layout_telemetry,
)

failures: list[str] = []
passed_checks = 0

TITLE_Y = 280_000
ONE_LINE_BAND = 700_000
TWO_LINE_BAND = 950_000
OLD_LIST_AVAIL = 4_865_200


def check(name: str, ok: bool, detail: str = "") -> None:
    global passed_checks
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if ok:
        passed_checks += 1
    else:
        failures.append(name)


def fresh(page: int = 4, key: str = "p04_density") -> tuple[Any, _Ctx]:
    reset_layout_telemetry()
    reset_bullet_measure()
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    return prs, _Ctx(prs, page, 80, slide_key=key)


def p1_title_band_by_fact() -> None:
    _, ctx = fresh()
    one = ctx.title("Резюме")
    check("П1а: однострочный заголовок отдаёт y + 700 000", one == TITLE_Y + ONE_LINE_BAND, f"получено {one}")
    _, ctx2 = fresh()
    long_title = "Россия — снимок выдачи: почему выделено (продолжение 2/3) — " + "очень длинная подпись " * 4
    two = ctx2.title(long_title)
    check("П1б: двухстрочный заголовок отдаёт y + 950 000", two == TITLE_Y + TWO_LINE_BAND, f"получено {two}")


def p5_title_box_ends_inside_its_band() -> None:
    """Рамка заголовка кончается внутри своей полосы: содержимое от `y + полоса` её не задевает.

    Ворота приёмки считают пересечения фигур по телеметрии: полоса стала
    короче, а рамка заголовка осталась 900 000 — и каждая страница с
    однострочным заголовком получила пересечение рамки с содержимым
    (`geometry: overlaps=44` на эталоне-72).
    """
    prs, ctx = fresh()
    next_y = ctx.title("Резюме")
    box = next(
        (sh for sh in prs.slides[0].shapes if getattr(sh, "has_text_frame", False) and "Резюме" in sh.text_frame.text),
        None,
    )
    bottom = int(box.top) + int(box.height) if box is not None else -1
    check(
        "П5: рамка однострочного заголовка кончается не ниже начала содержимого",
        box is not None and bottom <= next_y,
        f"низ рамки заголовка {bottom}, содержимое от {next_y}",
    )


def p2_bottom_reserve() -> None:
    check(
        "П2: CONTENT_BOTTOM == SLIDE_H − 800 000, подвал ниже",
        CONTENT_BOTTOM == SLIDE_H - 800_000 and FOOTER_Y > CONTENT_BOTTOM,
        f"CONTENT_BOTTOM = SLIDE_H − {SLIDE_H - CONTENT_BOTTOM}, FOOTER_Y = SLIDE_H − {SLIDE_H - FOOTER_Y}",
    )


def p3_p4_list_gets_the_room() -> None:
    prs, ctx = fresh()
    y = ctx.title("Резюме")
    block = "\n".join(["Тема блока", "Первая строка блока: 3 материала.", "Вторая строка блока."])
    ctx.bullets([block] * 3, y, max_items=9, max_chars=900)
    pages = [p for p in get_bullet_measure() if p.get("slideKey") == ctx.slide_key]
    avail = int(pages[0]["availableHeight"]) if pages else -1
    check(
        "П3: список под однострочным заголовком получает на 550 000 больше прежнего листа",
        avail == OLD_LIST_AVAIL + 550_000,
        f"availableHeight = {avail} (прежний лист {OLD_LIST_AVAIL})",
    )
    box = next(
        (
            sh
            for sh in prs.slides[0].shapes
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.startswith(BULLET_GLYPH)
        ),
        None,
    )
    bottom = int(box.top) + int(box.height) if box is not None else -1
    check(
        "П4: низ рамки списка не ниже CONTENT_BOTTOM, подвал ниже содержимого",
        box is not None and bottom <= CONTENT_BOTTOM < FOOTER_Y,
        f"низ рамки {bottom}, CONTENT_BOTTOM {CONTENT_BOTTOM}, FOOTER_Y {FOOTER_Y}",
    )


def main() -> int:
    p1_title_band_by_fact()
    p5_title_box_ends_inside_its_band()
    p2_bottom_reserve()
    p3_p4_list_gets_the_room()
    print(f"\n{'FAILED (' + str(len(failures)) + ')' if failures else 'PASSED (0 failures)'}")
    print_tap_counters(passed=passed_checks, failed=len(failures))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
