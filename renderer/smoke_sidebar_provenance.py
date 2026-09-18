#!/usr/bin/env python3
"""Смок: подпись происхождения боковой панели лежит внутри панели (шаг 0103).

Подпись в две строки выходила за нижний край подложки, а при блоках, кончающихся
у самого низа, рисовалась поверх последнего блока: рамка была фиксированной
высоты в одну строку, а верх — «не ниже `max_bottom − 120 000`». Теперь высота
подписи меряется, запас под неё у блоков — измеренный, и подпись либо стоит
после последнего блока целиком, либо не рисуется, а потеря названа.

Геометрия «внутри подложки» судится чужим прибором — вёрсткой LibreOffice.
Остальное — по фигурам PPTX. Сеть и база не нужны.

Запуск: python3 renderer/smoke_sidebar_provenance.py
"""

from __future__ import annotations

import base64
import io
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent))

import fitz  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

from smoke_counters import print_tap_counters  # noqa: E402
from orion_golden_render import render_orion_golden  # noqa: E402
from orion_golden_render.common import (  # noqa: E402
    FS_CAPTION,
    SLIDE_H,
    SLIDE_W,
    _Ctx,
    get_layout_telemetry,
    measure_text_height,
    reset_layout_telemetry,
)
from orion_golden_render.visual import _sidebar_analysis  # noqa: E402

failures: list[str] = []
passed_checks = 0
EMU_PER_PT = 12_700

PROVENANCE_ONE = "Источники на снимке — x.com и rupep.org."
PROVENANCE_TWO = "Источник — поисковая выдача: у показанных элементов нет отдельных адресов."
ANALYSIS = {
    "headlineConclusion": "На снимке выделено результатов повышенного внимания: 2.",
    "whatIsVisible": "На первой странице выдачи 10 результатов, из них 2 ведут на материалы с негативным контекстом.",
    "clientMeaning": "Выделенные материалы (2) видны при первичной проверке субъекта в этом регионе.",
    "recommendedActions": ["Проверить статусы 3 дел и первоисточники до принятия решений."],
}
#: Панель страницы подсказок золотого кейса (субъект синтетический): блоки
#: заполняют колонку до низа, и подпись в две строки на ней и вылезала.
TIGHT_ANALYSIS = {
    "sidebarMode": "adverse_explanation",
    "headlineConclusion": "Собрано 36, на панели — 10 подсказок: 8 относятся к субъекту и 2 требуют уточнения принадлежности. С негативной формулировкой — 5.",
    "highlightExplanations": [
        {"clientReason": "«Anders Holmström Nordkap Capital fraud»: отнесено к теме «Потенциально негативные публикации» по формулировке — требует ручной проверки.", "frameTone": "red"},
        {"clientReason": "«Anders Holmström Malta offshore»: отнесено к теме «Потенциально негативные публикации» по формулировке — требует ручной проверки.", "frameTone": "red"},
    ],
    "moreSignalsCount": 3,
    "clientMeaning": "Негативные подсказки видны пользователю ещё до просмотра результатов: они формируют первое впечатление и подталкивают к поиску компрометирующих материалов.",
    "recommendedActions": ["Сверить заголовки и домены с профилем субъекта; уточнить принадлежность строк со статусом «Вероятно»."],
}
#: Геометрия панели с подписью в одну строку до шага: верхи блоков и подписи.
TOPS_BEFORE = [1_470_000, 1_920_630, 2_120_630, 2_769_075, 2_969_075, 3_419_705, 3_619_705, 4_070_335]
SIDEBAR = dict(x=7_600_000, y=1_400_000, w=3_900_000, h=4_400_000)


def check(name: str, ok: bool, detail: str = "") -> None:
    global passed_checks
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))
    if ok:
        passed_checks += 1
    else:
        failures.append(name)


def draw_sidebar(analysis: dict[str, Any], **geometry: int) -> Any:
    reset_layout_telemetry()
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    ctx = _Ctx(prs, 20, 66, slide_key="p10_ru_serp_visual")
    _sidebar_analysis(ctx, {"visualAnalysis": analysis}, **{**SIDEBAR, **geometry})
    return prs.slides[0]


def text_shapes(slide: Any) -> list[Any]:
    return [sh for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]


def provenance_shape(slide: Any, text: str) -> Any:
    found = [sh for sh in text_shapes(slide) if sh.text_frame.text.strip() == text]
    return found[0] if found else None


def blocks_of(slide: Any, provenance: str) -> list[Any]:
    """Блоки панели — все текстовые фигуры, кроме подписи (сравнение по тексту:
    python-pptx отдаёт новый прокси-объект на каждый обход, тождество не работает)."""
    return [sh for sh in text_shapes(slide) if sh.text_frame.text.strip() != provenance]


def label_needed(label: Any, text: str) -> int:
    """Сколько высоты подписи нужно на самом деле — мерой рендерера, не рамкой."""
    return measure_text_height(text, int(label.width), FS_CAPTION, line_spacing=1.2)


def panel_card(slide: Any, inside: Any = None) -> Any:
    """Подложка панели: наименьшая фигура без текста, внутри которой стоит `inside`;
    без ориентира — самая крупная фигура без текста (панель, нарисованная отдельно)."""
    cards = [sh for sh in slide.shapes if not (getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip())]
    if inside is not None:
        cards = [
            sh
            for sh in cards
            if int(sh.left) <= int(inside.left) <= int(sh.left) + int(sh.width)
            and int(sh.top) <= int(inside.top) <= int(sh.top) + int(sh.height)
        ]
        return min(cards, key=lambda sh: int(sh.width) * int(sh.height)) if cards else None
    return max(cards, key=lambda sh: int(sh.width) * int(sh.height)) if cards else None


def c1_two_lines_inside_by_libreoffice() -> None:
    # Страница подсказок: панель-картинка 1200×720, как её строит приложение;
    # боковая колонка берёт высоту от неё — на такой странице подпись и вылезала.
    img = Image.new("RGB", (1200, 720), (246, 248, 250))
    d = ImageDraw.Draw(img)
    for i in range(10):
        d.rectangle([40, 60 + i * 62, 1160, 110 + i * 62], outline=(210, 214, 220), width=2)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    payload = {
        "reportSpec": {"subject": {"displayName": "Субъект Проверки"}},
        "deckManifest": {
            "finalSlides": [
                {
                    "slideKey": "p11_ru_suggestions_yandex",
                    "template": "orion_golden_surface_panel",
                    "templateId": "suggestions",
                    "title": "Россия — подсказки Яндекса: 5 негативных формулировок",
                    "pageNumber": 1,
                    "totalPageCount": 1,
                    "assetRefs": ["panel1"],
                    "visualAnalysis": {**TIGHT_ANALYSIS, "provenanceLabel": PROVENANCE_TWO},
                }
            ]
        },
        "assets": [
            {
                "assetRef": "panel1",
                "kind": "surface_panel",
                "title": "Россия — поисковые подсказки Яндекса",
                "imageData": base64.b64encode(buf.getvalue()).decode("ascii"),
                "mimeType": "image/png",
            }
        ],
    }
    out = render_orion_golden(payload)
    if out.get("pdfExportMode") != "libreoffice":
        print("# SKIP С1 подпись внутри панели — LibreOffice недоступен, вёрстку судить нечем")
        return
    prs = Presentation(io.BytesIO(base64.b64decode(out["pptxBase64"])))
    slide = prs.slides[0]
    label = provenance_shape(slide, PROVENANCE_TWO)
    card = panel_card(slide, label) if label is not None else None
    if card is None or label is None:
        check("С1: подпись и подложка панели найдены на странице", False, f"подложка {card is not None}, подпись {label is not None}")
        return
    doc = fitz.open(stream=base64.b64decode(out["pdfBase64"]), filetype="pdf")
    left_pt, right_pt = int(card.left) / EMU_PER_PT, (int(card.left) + int(card.width)) / EMU_PER_PT
    card_bottom_pt = (int(card.top) + int(card.height)) / EMU_PER_PT
    words = set(PROVENANCE_TWO.split())
    bottoms = [
        span["bbox"][3]
        for block in doc[0].get_text("dict")["blocks"]
        for line in block.get("lines", [])
        for span in line.get("spans", [])
        if span["text"].strip()
        and left_pt - 1 <= span["bbox"][0] <= right_pt
        and int(label.top) / EMU_PER_PT - 2 <= span["bbox"][1]
        and words & set(span["text"].split())
    ]
    check("С1а: строки подписи найдены в PDF", len(bottoms) >= 2, f"строк {len(bottoms)}")
    check(
        "С1б: последняя строка подписи не ниже низа подложки панели (LibreOffice)",
        bool(bottoms) and max(bottoms) <= card_bottom_pt + 0.5,
        f"низ подписи {max(bottoms):.1f} pt при низе подложки {card_bottom_pt:.1f} pt" if bottoms else "строк нет",
    )


def c2_one_line_keeps_geometry() -> None:
    slide = draw_sidebar({**ANALYSIS, "provenanceLabel": PROVENANCE_ONE})
    shapes = text_shapes(slide)
    tops = [int(sh.top) for sh in shapes]
    check("С2а: подпись в одну строку — верхи блоков панели прежние", tops == TOPS_BEFORE, f"верхи {tops}")
    label = provenance_shape(slide, PROVENANCE_ONE)
    blocks = blocks_of(slide, PROVENANCE_ONE)
    last_bottom = max(int(sh.top) + int(sh.height) for sh in blocks)
    check(
        "С2б: подпись стоит после последнего блока",
        label is not None and int(label.top) >= last_bottom,
        f"верх подписи {int(label.top) if label else None}, низ последнего блока {last_bottom}",
    )


def c3_c5_tight_column() -> None:
    long_visible = (
        "На первой странице выдачи 10 результатов, из них 2 ведут на материалы с негативным контекстом; "
        "остальные 8 — деловые и биографические публикации, которые не меняют оценку. "
        "Порядок выдачи снят в момент прогона и может измениться при следующем запросе."
    )
    # Колонка, в которой блоки кончаются у самого низа: подпись не вправе лечь
    # поверх последнего блока и не вправе выйти за подложку.
    slide = draw_sidebar({**ANALYSIS, "whatIsVisible": long_visible, "provenanceLabel": PROVENANCE_TWO}, h=3_000_000)
    shapes = text_shapes(slide)
    label = provenance_shape(slide, PROVENANCE_TWO)
    blocks = blocks_of(slide, PROVENANCE_TWO)
    last_bottom = max(int(sh.top) + int(sh.height) for sh in blocks)
    card = panel_card(slide)
    card_bottom = int(card.top) + int(card.height) if card is not None else 0
    losses = [e for e in get_layout_telemetry() if e.get("name", "").startswith("orion_sidebar_provenanceLabel")]
    if label is not None:
        needed = label_needed(label, PROVENANCE_TWO)
        check(
            "С3: подпись не ложится поверх последнего блока и не выходит за подложку",
            int(label.top) >= last_bottom and int(label.top) + needed <= card_bottom,
            f"верх подписи {int(label.top)}, низ последнего блока {last_bottom}, низ текста по мере {int(label.top) + needed}, низ подложки {card_bottom}",
        )
        # Рамка — в свою меру: инспектор геометрии судит по рамкам, и рамка в
        # одну строку под текстом в две прятала бы от него настоящий низ.
        check(
            "С3б: рамка подписи не ниже своей меры — инспектор геометрии видит настоящий низ",
            int(label.height) >= needed,
            f"рамка {int(label.height)}, мера {needed}",
        )
    else:
        check(
            "С4: места нет — подпись не нарисована, потеря названа предупреждением, блоки целы",
            len(losses) == 1 and int(losses[0].get("droppedLines") or 0) == 0 and len(blocks) >= 6,
            f"записей о подписи {len(losses)}, блоков {len(blocks)}",
        )
    # Тесная колонка: обязательный вывод и «Что сделать» напечатаны целиком,
    # а подпись — только целиком или никак.
    headline = [sh for sh in shapes if sh.text_frame.text.strip() == ANALYSIS["headlineConclusion"]]
    action = [sh for sh in shapes if sh.text_frame.text.strip() == ANALYSIS["recommendedActions"][0]]
    check(
        "С5: на тесной колонке вывод и «Что сделать» напечатаны целиком",
        len(headline) == 1 and len(action) == 1,
        f"вывод {len(headline)}, действие {len(action)}",
    )
    # Совсем нет места: колонка в 560 000 EMU, где даже обязательный вывод
    # заменяется запасной фразой. Подпись обязана исчезнуть с записью, а не
    # вылезти за подложку.
    slide = draw_sidebar({**ANALYSIS, "provenanceLabel": PROVENANCE_TWO}, h=560_000)
    label = provenance_shape(slide, PROVENANCE_TWO)
    card = panel_card(slide)
    card_bottom = int(card.top) + int(card.height) if card is not None else 0
    losses = [e for e in get_layout_telemetry() if e.get("name", "").startswith("orion_sidebar_provenanceLabel")]
    text_bottom = int(label.top) + label_needed(label, PROVENANCE_TWO) if label is not None else None
    inside = label is None or text_bottom <= card_bottom
    check(
        "С4: места нет — подпись либо целиком внутри, либо не нарисована с записью-предупреждением",
        inside and (label is not None or (len(losses) == 1 and int(losses[0].get("droppedLines") or 0) == 0)),
        f"подпись {'есть' if label is not None else 'нет'}, низ текста по мере {text_bottom}, низ подложки {card_bottom}, записей {len(losses)}",
    )


def main() -> int:
    c1_two_lines_inside_by_libreoffice()
    c2_one_line_keeps_geometry()
    c3_c5_tight_column()
    print(f"\n{'FAILED (' + str(len(failures)) + ')' if failures else 'PASSED (0 failures)'}")
    print_tap_counters(passed=passed_checks, failed=len(failures))
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
